#!/usr/bin/env python3
"""Export data/merged/timeline.json (+ series.json) to a review workbook and a flat CSV.

  exports/timeline_review.xlsx  – Read me · Timeline (one row per record, with review columns) · Series ·
                                  Series members · Add new events (template) · Sources · Categories · Summary
  exports/timeline.csv          – the Timeline sheet as UTF-8 CSV (with BOM so Excel shows Tamil)

Regenerate after every merge:  python3 scripts/export_review.py
"""
import csv, json, os
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = json.load(open(os.path.join(ROOT, "data/merged/timeline.json"), encoding="utf-8"))
SCHEMA = json.load(open(os.path.join(ROOT, "schema/event.schema.json"), encoding="utf-8"))
EVENTS = DATA["events"]
SERIES = json.load(open(os.path.join(ROOT, "data/merged/series.json"), encoding="utf-8"))["series"]
BY_ID = {e["id"]: e for e in EVENTS}
HINT = {"covered_by_span": "Covered by span – hide on chart", "span_endpoint": "Span endpoint – keep",
        "index_entry": "Index list entry – app detail"}
OUT_X = os.path.join(ROOT, "exports/timeline_review.xlsx")
OUT_C = os.path.join(ROOT, "exports/timeline.csv")

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")   # owner fills these
BAND_FILL = {"bible": "DDEBF7", "history": "FCE4D6", "jw": "E2EFDA", "future": "E4DFEC"}
THIN = Side(style="thin", color="BFBFBF")
PREC = {"exact": "Exact", "approx": "Approx. (c.)", "before": "Before (b.)", "after": "After (a.)",
        "between": "Between (one moment in a window)", "unknown": "Not stated", "ongoing": "Ongoing", "undated": "Future (undated)"}

BAND_NAME = {"bible": "Bible", "history": "Bible & church history", "jw": "Jehovah's Witnesses", "future": "Future"}
CATEGORY_HELP = {
    "bible.creation": "Creation", "bible.patriarch": "Patriarchs (Adam to Joseph)", "bible.covenant": "Covenants",
    "bible.judgment": "Divine judgments", "bible.exodus_law": "Exodus and the Law", "bible.conquest_judges": "Conquest and Judges",
    "bible.monarchy": "Kings of Israel and Judah", "bible.prophet": "Prophets", "bible.temple": "Temple",
    "bible.exile_return": "Exile and return", "bible.book_writing": "Bible books written", "bible.intertestament": "Between the Testaments",
    "bible.jesus": "Jesus' life and ministry", "bible.apostolic": "Apostles and early congregation",
    "bible.prophecy_fulfilled": "Prophecy fulfilled", "history.apostasy": "Apostasy and Christendom",
    "history.bible_translation": "Bible translation", "history.bible_preservation": "Bible preservation (manuscripts, canon)",
    "history.reform": "Reformers and reform movements", "history.persecution": "Persecution of Bible readers",
    "jw.organization": "Organization", "jw.doctrine": "Beliefs clarified", "jw.publication": "Publications",
    "jw.legal": "Legal cases", "jw.persecution": "Persecution and bans", "jw.convention": "Conventions",
    "jw.construction": "Construction (Bethel, halls)", "jw.preaching": "Preaching work", "jw.leadership": "Leadership",
    "jw.education": "Theocratic schools", "jw.relief": "Disaster relief", "future.prophecy": "Future prophetic events",
}


def year_txt(p):
    if not p or p.get("year") is None:
        return ""
    return f"{p['year']} {'BCE' if p['era'] == 'BCE' else 'CE'}"


def signed(p):
    if not p or p.get("year") is None:
        return None
    return -p["year"] if p["era"] == "BCE" else p["year"]


def refs(ev, key):
    return "; ".join(s.get(key) or "" for s in ev.get("scriptures", []) if s.get(key))


# ---- column model: (header, width, getter, kind) kind: text|num|link|review
def primary(ev):
    return ev["sources"][0] if ev["sources"] else {}


COLUMNS = [
    ("ID", 10, lambda e: e["id"], "text"),
    ("Type", 7, lambda e: "Span" if e["timing"] == "span" else "Point", "text"),
    ("Span kind", 15, lambda e: (e.get("span_kind") or "").replace("_", " "), "text"),
    ("Date (as printed)", 26, lambda e: e["date_label_en"], "text"),
    ("Date (Tamil)", 22, lambda e: e.get("date_label_ta") or "", "text"),
    ("Start year", 10, lambda e: year_txt(e["start"]), "text"),
    ("Start precision", 13, lambda e: PREC.get(e["start"]["precision"], e["start"]["precision"]), "text"),
    ("End year", 10, lambda e: year_txt(e.get("end")) if e.get("end") else "", "text"),
    ("End precision", 13, lambda e: PREC.get(e["end"]["precision"], "") if e.get("end") else "", "text"),
    ("Duration", 13, lambda e: (e.get("duration") or {}).get("display_en") or "", "text"),
    ("Title (English)", 46, lambda e: e["title_en"], "text"),
    ("Title (Tamil)", 40, lambda e: e.get("title_ta") or "", "text"),
    ("Description (English)", 60, lambda e: e.get("description_en") or "", "text"),
    ("Description (Tamil)", 50, lambda e: e.get("description_ta") or "", "text"),
    ("Band", 14, lambda e: BAND_NAME[e["category"].split(".")[0]], "text"),
    ("Category", 22, lambda e: e["category"], "text"),
    ("Event kind", 14, lambda e: e.get("event_kind") or "", "text"),
    ("People", 22, lambda e: "; ".join(e.get("people") or []), "text"),
    ("Places", 20, lambda e: "; ".join(e.get("places") or []), "text"),
    ("Scriptures (English)", 22, lambda e: refs(e, "ref_en"), "text"),
    ("Scriptures (Tamil)", 20, lambda e: refs(e, "ref_ta"), "text"),
    ("Sources", 8, lambda e: len(e["sources"]), "num"),
    ("Main source", 34, lambda e: primary(e).get("title_en") or "", "text"),
    ("Main source link (EN)", 16, lambda e: primary(e).get("url_en") or "", "link"),
    ("Main source link (TA)", 16, lambda e: primary(e).get("url_ta") or "", "link"),
    ("All source links (EN)", 30, lambda e: "\n".join(s["url_en"] for s in e["sources"]), "text"),
    ("Images", 7, lambda e: len(e.get("images") or []), "num"),
    ("First image", 16, lambda e: (e.get("images") or [{}])[0].get("url", "") if e.get("images") else "", "link"),
    ("Sources differ", 8, lambda e: len(e.get("conflicts") or []), "num"),
    ("Where sources differ", 40, lambda e: " | ".join(c["note"] for c in (e.get("conflicts") or [])), "text"),
    ("Computed by pipeline", 36, lambda e: (e.get("derived") or {}).get("note", ""), "text"),
    ("Start event ID", 10, lambda e: e.get("start_event_id") or "", "text"),
    ("End event ID", 10, lambda e: e.get("end_event_id") or "", "text"),
    ("Inside span", 10, lambda e: e.get("parent_id") or "", "text"),
    ("Tier", 6, lambda e: e.get("tier") or "", "num"),
    ("On A0 poster", 8, lambda e: "Yes" if e.get("poster") else "", "text"),
    ("Tier reason", 30, lambda e: e.get("tier_why") or "", "text"),
    ("Chart hint", 24, lambda e: HINT.get(e.get("chart_hint"), ""), "text"),
    ("Start/end of span", 34, lambda e: "; ".join(f"{r['as']} of {r['span_title']} ({r['span_id']})" for r in e.get("endpoint_of") or []), "text"),
    ("Series", 30, lambda e: "; ".join(f"{m['id']} #{m['order']} ({m['role']})" for m in e.get("series") or []), "text"),
    ("Notes", 30, lambda e: e.get("notes") or "", "text"),
    ("Sort key", 9, lambda e: e["start"]["sort_key"], "num"),
    # review columns, filled by the owner
    ("✎ Status", 12, lambda e: "", "review"),
    ("✎ Importance (1-3)", 11, lambda e: e.get("importance") or "", "review"),
    ("✎ Show on chart", 10, lambda e: "", "review"),
    ("✎ Correction / comment", 40, lambda e: "", "review"),
]
COL = {h: i + 1 for i, (h, *_rest) in enumerate(COLUMNS)}


def head(ws, row, headers, widths=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.row_dimensions[row].height = 32


def style_body(ws, max_row, max_col):
    for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
        for cell in row:
            if cell.font is None or cell.font.name != FONT:
                cell.font = Font(name=FONT, size=10, color=cell.font.color if cell.font else None,
                                 underline=cell.font.underline if cell.font else None)


def build_timeline(wb):
    ws = wb.create_sheet("Timeline")
    head(ws, 1, [c[0] for c in COLUMNS], [c[1] for c in COLUMNS])
    for r, ev in enumerate(EVENTS, 2):
        band = ev["category"].split(".")[0]
        for c, (h, _w, get, kind) in enumerate(COLUMNS, 1):
            v = get(ev)
            cell = ws.cell(row=r, column=c, value=v if v != "" else None)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=h in ("Title (English)", "Title (Tamil)"))
            if kind == "link" and v:
                cell.hyperlink = v
                cell.value = "open"
                cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")
            elif kind == "review":
                cell.fill = REVIEW_FILL
                cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            elif h == "Band":
                cell.fill = PatternFill("solid", fgColor=BAND_FILL[band])
            elif h == "Type" and v == "Span":
                cell.font = Font(name=FONT, size=10, bold=True)
    last = len(EVENTS) + 1
    ws.freeze_panes = ws.cell(row=2, column=COL["Title (English)"])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    status = DataValidation(type="list", formula1='"OK,Fix,Remove,Add Tamil,Duplicate"', allow_blank=True)
    imp = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    show = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(status)
    ws.add_data_validation(imp)
    ws.add_data_validation(show)
    show.add(f"{get_column_letter(COL['✎ Show on chart'])}2:{get_column_letter(COL['✎ Show on chart'])}{last}")
    status.add(f"{get_column_letter(COL['✎ Status'])}2:{get_column_letter(COL['✎ Status'])}{last}")
    imp.add(f"{get_column_letter(COL['✎ Importance (1-3)'])}2:{get_column_letter(COL['✎ Importance (1-3)'])}{last}")
    # hide technical columns by default (still filterable)
    for h in ("Sort key", "All source links (EN)", "Start event ID", "End event ID", "Inside span", "Notes"):
        ws.column_dimensions[get_column_letter(COL[h])].hidden = True
    return ws


def build_series(wb):
    ws = wb.create_sheet("Series")
    hdr = ["Series ID", "Title", "Kind", "Highlight on chart", "Auto-built", "Members", "First", "Last", "Description",
           "Main source", "Source link", "✎ Status", "✎ Comment"]
    head(ws, 1, hdr, [26, 40, 18, 10, 9, 9, 22, 22, 70, 36, 12, 12, 40])
    for r, sd in enumerate(sorted(SERIES, key=lambda x: (x.get("auto", False), x["id"])), 2):
        mem = [BY_ID[m["event_id"]] for m in sd["members"] if m["event_id"] in BY_ID]
        dated = sorted(mem, key=lambda e: e["start"]["sort_key"])
        src = sd["sources"][0] if sd.get("sources") else {}
        vals = [sd["id"], sd["title_en"], sd["kind"].replace("_", " "), "Yes" if sd.get("highlight") else "",
                "Yes" if sd.get("auto") else "", len(mem), dated[0]["date_label_en"] if dated else "",
                dated[-1]["date_label_en"] if dated else "", sd["description_en"], src.get("title_en"), src.get("url_en"), "", ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v if v != "" else None)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=c in (2, 9))
            if c == 11 and v:
                cell.hyperlink, cell.value = v, "open"
                cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")
            if c >= 12:
                cell.fill = REVIEW_FILL
                cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:M{len(SERIES) + 1}"

    wm = wb.create_sheet("Series members")
    hdr = ["Series ID", "Order", "Role", "Group", "Caption in series", "Record ID", "Type", "Date (as printed)",
           "Record title", "Pairs with", "Scriptures for this step", "Note", "✎ Comment"]
    head(wm, 1, hdr, [26, 7, 16, 18, 46, 10, 7, 24, 46, 22, 22, 36, 36])
    r = 2
    for sd in sorted(SERIES, key=lambda x: (x.get("auto", False), x["id"])):
        groups = {g["id"]: g["label_en"] for g in sd.get("groups", [])}
        for m in sorted(sd["members"], key=lambda m: m["order"]):
            e = BY_ID.get(m["event_id"])
            if not e:
                continue
            vals = [sd["id"], m["order"], m["role"].replace("_", " "), groups.get(m.get("group"), m.get("group") or ""),
                    m.get("label_en") or "", e["id"], "Span" if e["timing"] == "span" else "Point", e["date_label_en"],
                    e["title_en"], ", ".join(m.get("pairs_with") or []), "; ".join(m.get("scriptures") or []), m.get("note") or "", ""]
            for c, v in enumerate(vals, 1):
                cell = wm.cell(row=r, column=c, value=v if v != "" else None)
                cell.font = Font(name=FONT, size=10)
                if c == 13:
                    cell.fill = REVIEW_FILL
            r += 1
    wm.freeze_panes = "C2"
    wm.auto_filter.ref = f"A1:M{r - 1}"


def build_add(wb):
    ws = wb.create_sheet("Add new events")
    hdr = ["Type (Point/Span)", "Date as printed (English)", "Date as printed (Tamil)", "Start year", "Start era (BCE/CE)",
           "Start precision", "End year (spans)", "End era", "End precision", "Duration as printed", "Title (English)",
           "Title (Tamil)", "Description (English)", "Description (Tamil)", "Category", "People (; separated)",
           "Places (; separated)", "Scriptures (English)", "Scriptures (Tamil)", "jw.org / wol.jw.org link (required)",
           "Tamil jw.org link", "Image link", "Your note"]
    widths = [12, 24, 22, 10, 10, 14, 10, 9, 14, 14, 40, 36, 50, 44, 22, 22, 20, 22, 20, 44, 40, 30, 30]
    head(ws, 1, hdr, widths)
    example = ["Point", "455 B.C.E., Elul 25", "கி.மு. 455, எலூல் 25", 455, "BCE", "exact", None, None, None, "52 days",
               "EXAMPLE – Jerusalem's wall completed in 52 days", "", "Under Nehemiah's direction the wall is finished.",
               "", "bible.exile_return", "Nehemiah", "Jerusalem", "Neh. 6:15", "நெ. 6:15",
               "https://wol.jw.org/en/wol/d/r1/lp-e/1101990130", "https://wol.jw.org/ta/wol/d/r122/lp-tl/1101990130", "",
               "Example row – overwrite or delete"]
    for c, v in enumerate(example, 1):
        cell = ws.cell(row=2, column=c, value=v)
        cell.font = Font(name=FONT, size=10, italic=True, color="7F7F7F")
    for r in range(3, 203):
        for c in range(1, len(hdr) + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = REVIEW_FILL
            cell.font = Font(name=FONT, size=10)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.freeze_panes = "B2"
    dv = lambda f: DataValidation(type="list", formula1=f, allow_blank=True)
    v_type, v_era = dv('"Point,Span"'), dv('"BCE,CE"')
    v_prec = dv('"exact,approx,before,after,between,unknown,ongoing,undated"')
    v_cat = dv("=Categories!$A$2:$A$%d" % (1 + len(CATEGORY_HELP)))
    for v, cols in ((v_type, "A"), (v_era, "E"), (v_era, "H"), (v_prec, "F"), (v_prec, "I"), (v_cat, "O")):
        ws.add_data_validation(v)
        v.add(f"{cols}2:{cols}202")


def build_sources(wb):
    ws = wb.create_sheet("Sources")
    agg = {}
    for ev in EVENTS:
        for s in ev["sources"]:
            a = agg.setdefault(s["docid"], {"title": s["title_en"], "pub": s.get("publication") or "", "en": s["url_en"],
                                            "ta": s.get("url_ta") or "", "n": 0})
            a["n"] += 1
            if not a["ta"] and s.get("url_ta"):
                a["ta"] = s["url_ta"]
    rows = sorted(agg.items(), key=lambda kv: -kv[1]["n"])
    head(ws, 1, ["Document ID", "Title", "Publication", "Records citing it", "English link", "Tamil link"], [14, 60, 34, 12, 14, 14])
    for r, (docid, a) in enumerate(rows, 2):
        vals = [docid, a["title"], a["pub"], a["n"], a["en"], a["ta"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v or None)
            cell.font = Font(name=FONT, size=10)
            if c in (5, 6) and v:
                cell.hyperlink, cell.value = v, "open"
                cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:F{len(rows) + 1}"
    return len(rows)


def build_categories(wb):
    ws = wb.create_sheet("Categories")
    head(ws, 1, ["Category code", "Band", "Meaning", "Records"], [26, 22, 40, 10])
    cat_col = get_column_letter(COL["Category"])
    for r, (code, meaning) in enumerate(CATEGORY_HELP.items(), 2):
        band = code.split(".")[0]
        for c, v in enumerate([code, BAND_NAME[band], meaning], 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            if c == 2:
                cell.fill = PatternFill("solid", fgColor=BAND_FILL[band])
        cell = ws.cell(row=r, column=4, value=f"=COUNTIF(Timeline!${cat_col}:${cat_col},A{r})")
        cell.font = Font(name=FONT, size=10)
    total_r = len(CATEGORY_HELP) + 2
    ws.cell(row=total_r, column=3, value="Total").font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=total_r, column=4, value=f"=SUM(D2:D{total_r - 1})").font = Font(name=FONT, size=10, bold=True)


def build_summary(wb):
    ws = wb.create_sheet("Summary")
    L = lambda h: get_column_letter(COL[h])
    t, typ, band, tta, imgs, diff, sp = (L("ID"), L("Type"), L("Band"), L("Title (Tamil)"), L("Images"),
                                          L("Sources differ"), L("Start precision"))
    rows = [
        ("Records", f"=COUNTA(Timeline!{t}:{t})-1"),
        ("Points (one moment)", f'=COUNTIF(Timeline!{typ}:{typ},"Point")'),
        ("Spans (start to end)", f'=COUNTIF(Timeline!{typ}:{typ},"Span")'),
        ("Bible", f'=COUNTIF(Timeline!{band}:{band},"Bible")'),
        ("Bible & church history", f'=COUNTIF(Timeline!{band}:{band},"Bible & church history")'),
        ("Jehovah's Witnesses", f'=COUNTIF(Timeline!{band}:{band},"Jehovah\'s Witnesses")'),
        ("Future", f'=COUNTIF(Timeline!{band}:{band},"Future")'),
        ("With Tamil title", f"=COUNTA(Timeline!{tta}:{tta})-1"),
        ("With images", f'=COUNTIF(Timeline!{imgs}:{imgs},">0")'),
        ("Where sources differ", f'=COUNTIF(Timeline!{diff}:{diff},">0")'),
        ("Uncertain moments (between)", f'=COUNTIF(Timeline!{sp}:{sp},"Between*")'),
        ("Marked OK by you", f'=COUNTIF(Timeline!{L("✎ Status")}:{L("✎ Status")},"OK")'),
        ("Marked Fix by you", f'=COUNTIF(Timeline!{L("✎ Status")}:{L("✎ Status")},"Fix")'),
        ("Marked Remove by you", f'=COUNTIF(Timeline!{L("✎ Status")}:{L("✎ Status")},"Remove")'),
        ("New events you added", "=COUNTA('Add new events'!K3:K202)"),
    ]
    head(ws, 1, ["Measure", "Count"], [34, 12])
    for r, (k, f) in enumerate(rows, 2):
        ws.cell(row=r, column=1, value=k).font = Font(name=FONT, size=10)
        ws.cell(row=r, column=2, value=f).font = Font(name=FONT, size=10)
    ws.cell(row=len(rows) + 3, column=1, value="Counts are live formulas over the Timeline and Add new events sheets.").font = Font(name=FONT, size=9, italic=True, color="7F7F7F")


def build_readme(wb, n_sources):
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    lines = [
        ("JW Timeline – data review", None),
        ("Data version", f"{DATA['generated']} · {len(EVENTS)} records · schema v{DATA['schema_version']} · github.com/jeyantx/jw-timeline"),
        ("Sources", f"Every record cites wol.jw.org or www.jw.org only ({n_sources} documents; see the Sources sheet)."),
        ("", None),
        ("SHEETS", None),
        ("Timeline", "One row per record, oldest first. Use the filter arrows (e.g. Type = Span, Band = Bible, Category, Chart hint, Series). Hidden technical columns can be unhidden."),
        ("Series", "Curated groups that belong together and get their own colour on the chart: prophecy and fulfillment, world powers of Daniel, covenants, successions of kings, reproduced jw.org charts."),
        ("Series members", "Every member of every series in order, with its role (prophecy, fulfillment, rise, fall, ruler…) and caption."),
        ("Add new events", "Blank rows for anything missing. A jw.org / wol.jw.org link is required for every new row. Row 2 is a grey example."),
        ("Sources", "Every jw.org document used, with how many records cite it."),
        ("Categories", "The 32 categories, their meaning and live counts."),
        ("Summary", "Live counts, including how many rows you have marked."),
        ("", None),
        ("HOW TO REVIEW", None),
        ("Yellow cells", "Are yours to fill. Everything else is the data as collected – change it only through the comment column, so the change can be traced back to a source."),
        ("✎ Status", "OK · Fix (say what in the comment) · Remove · Add Tamil · Duplicate (name the other ID in the comment)."),
        ("✎ Importance (1-3)", "1 = must appear on the A0 chart · 2 = major · 3 = detail (app only). Leave blank if unsure."),
        ("✎ Correction / comment", "What is wrong and the jw.org link that shows the right value."),
        ("", None),
        ("READING THE COLUMNS", None),
        ("Type", "Point = one moment. Span = lasted from a start to an end (a life, a reign, a journey, a ban)."),
        ("Date (as printed)", "Exactly as the jw.org source prints it. c. = about, a. = after, b. = before."),
        ("Start / End precision", "Exact · Approx. (c.) · Before (b.) · After (a.) · Between (happened once, somewhere in a window, e.g. '31 or 32 C.E.') · Not stated (jw.org gives no year; only the order is known) · Ongoing (span not ended) · Future (undated)."),
        ("Duration", "As printed ('41 years'), or computed from the start and end."),
        ("Chart hint", "Covered by span – a birth, death, creation or accession that only marks the start/end of a life, reign or office bar; the chart can hide it and show the bar. Span endpoint – starts/ends a span but matters on its own (kept). Index list entry – one-line entry from the Watch Tower Publications Index lists (beliefs clarified, convention themes); app detail. Use '✎ Show on chart' to override."),
        ("Series", "Which series a record belongs to, with its position and role."),
        ("Computed by pipeline", "Filled when the record or a value was built from other sourced records – e.g. a lifespan from a birth and a death, a reign ending when the next king begins (NWT App. A6). Check these especially."),
        ("Sources differ", "Number of places where two jw.org publications give different values; the next column says what differs."),
        ("Tamil columns", "Only filled with text copied from Tamil jw.org pages. Blank = no Tamil page found (Insight and Proclaimers have no Tamil edition)."),
        ("Links", "'open' cells link to the page on wol.jw.org."),
        ("", None),
        ("SEND IT BACK", None),
        ("", "Save the file and share it back (or upload to the repo). The changes will be applied to the JSON and a new version exported."),
    ]
    for r, (a, b) in enumerate(lines, 1):
        ca = ws.cell(row=r, column=1, value=a or None)
        cb = ws.cell(row=r, column=2, value=b)
        ca.font = Font(name=FONT, size=11, bold=bool(a) and (b is None or a.isupper() or r == 1))
        cb.font = Font(name=FONT, size=10)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.alignment = Alignment(vertical="top")
        if a in ("Yellow cells", "✎ Status", "✎ Importance (1-3)", "✎ Correction / comment"):
            ca.fill = REVIEW_FILL
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color="1F3864")


def write_csv():
    with open(OUT_C, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        hdr = [c[0] for c in COLUMNS]
        w.writerow(hdr)
        for ev in EVENTS:
            w.writerow([("" if (v := get(ev)) is None else v) for (_h, _w, get, _k) in COLUMNS])


def main():
    os.makedirs(os.path.dirname(OUT_X), exist_ok=True)
    wb = Workbook()
    build_timeline(wb)
    build_series(wb)
    build_add(wb)
    n_sources = build_sources(wb)
    build_categories(wb)
    build_summary(wb)
    build_readme(wb, n_sources)
    wb.move_sheet("Read me", offset=-wb.index(wb["Read me"]))
    wb.active = 0
    wb.save(OUT_X)
    write_csv()
    print(f"wrote {OUT_X} ({len(EVENTS)} rows, {n_sources} sources) and {OUT_C}")


if __name__ == "__main__":
    main()
